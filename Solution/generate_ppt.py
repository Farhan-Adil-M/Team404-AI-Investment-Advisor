"""generate_ppt.py — builds Team404 submission deck (Problem → Approach → Solution → Demo → Impact).
Run: python generate_ppt.py  →  Team404_Deck.pptx"""
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor

CYAN, PINK, DARK, WHITE = RGBColor(0, 0xE5, 0xFF), RGBColor(0xFF, 0x2B, 0x85), RGBColor(0x0A, 0x06, 0x1D), RGBColor(0xFF, 0xFF, 0xFF)

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)


def slide(title, body_lines, subtitle=""):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg = s.background.fill; bg.solid(); bg.fore_color.rgb = DARK
    tb = s.shapes.add_textbox(Inches(0.6), Inches(0.4), Inches(12.1), Inches(1.2)).text_frame
    tb.text = title
    tb.paragraphs[0].font.size, tb.paragraphs[0].font.bold = Pt(36), True
    tb.paragraphs[0].font.color.rgb = CYAN
    if subtitle:
        p = tb.add_paragraph(); p.text = subtitle
        p.font.size, p.font.color.rgb = Pt(16), PINK
    body = s.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(11.9), Inches(5.4)).text_frame
    body.word_wrap = True
    for i, line in enumerate(body_lines):
        p = body.paragraphs[0] if i == 0 else body.add_paragraph()
        p.text = line
        p.font.size = Pt(20 if not line.startswith("  ") else 16)
        p.font.color.rgb = WHITE
        p.space_after = Pt(10)
    return s

# 1 — Title
s = prs.slides.add_slide(prs.slide_layouts[6])
bg = s.background.fill; bg.solid(); bg.fore_color.rgb = DARK
t = s.shapes.add_textbox(Inches(0.8), Inches(2.2), Inches(11.7), Inches(3)).text_frame
t.text = "AI-Powered Personal Investment Advisor"
t.paragraphs[0].font.size, t.paragraphs[0].font.bold = Pt(44), True
t.paragraphs[0].font.color.rgb = CYAN
for line, sz, col in [("Background Analysis & Risk-Based Planning · Matrix Hackathon 2026", 20, PINK),
                      ("Team404 — A.kruthika · Shaik Zeeshan · M. Farhan Adil · M. Noel", 22, WHITE),
                      ("DRKVSRIT × Data Science · Oct 1, 2026", 16, WHITE)]:
    p = t.add_paragraph(); p.text = line; p.font.size = Pt(sz); p.font.color.rgb = col

slide("Problem", [
    "• Investors' right decision depends on budget, experience, background & risk comfort",
    "• No tool considers all of it together — beginners get generic advice",
    "• People who took losses don't know WHY it happened or how to recover",
    "• Emotion (hype, panic) drives decisions instead of data",
])

slide("Approach", [
    "• 5-step wizard: Rich Profile → Loss Analysis → Risk → Plan → Ongoing Optimizer",
    "• Asks what real advice needs: age, income, expenses, dependents, goals, emergency fund",
    "• Every claim proven against REAL market data (yfinance, 12 live tickers)",
    "• Deterministic AI rule-engine — works offline, cannot crash on API outages",
    "• Advisory rules: 110−age glide path · goal buckets · emergency-fund gate ·",
    "  risk capacity vs willingness (binding = min)",
    "• Built & verified by an agent workflow under 18 strict hackathon rules",
])

slide("Solution", [
    "• Loss taxonomy detects: bought-at-peak, panic-sell, hype-chase, concentration…",
    "  with hard evidence (e.g. “entry within 0.8% of 90-day high”)",
    "• DUAL PLAN: Recommended-for-you (profile-driven) + Build-your-own (live sliders,",
    "  expected return vs worst-case trade-off, guard-rail warnings)",
    "• High-risk: aggressive picks w/ live momentum, volatility, drawdown + reasoning",
    "• Low-risk: SIP with 3/5/10y scenarios, seasonality chart, dip-window timing",
    "• Optimizer: dip → top-up amount, drift >5pp → rebalance, vol spike → hold",
])

slide("Demo", [
    "1. Rich intake: age 28, ₹80k income, goal 12y, risk willingness 4/5, budget ₹1L",
    "2. Loss: bought 2025-02-11 (peak!), exited 2025-03-13 (−8.8% dip) → AI explains",
    "3. Risk: YES → Recommended plan (76% equity, emergency-gate flag) vs Custom sliders",
    "4. Alternative: NO → SIP ₹10k/mo, 10y projection, best-month reminder",
    "5. Optimizer: live signals (“Crypto +32% → rebalance”)",
])

slide("Impact", [
    "• Beginners finally get explained, evidence-backed guidance — not generic tips",
    "• Turns past losses into specific, provable lessons instead of vague regret",
    "• Encourages disciplined SIP / rebalancing habits (Behavioural finance fix)",
    "• 100% free, open, offline-safe; works on any laptop",
    "• Judging: 30 pts live execution · 25 innovation · 20 UI · 15 impact · 10 fit",
])

prs.save("Team404_Deck.pptx")
print("saved Team404_Deck.pptx")
